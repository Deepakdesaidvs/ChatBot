"""Download formats for the conversation: Markdown, plain text and Word (.docx)."""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document

from src.memory import ConversationMemory, Turn

_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # characters Word files cannot contain
TITLE = "Document chat"


def _plain(text: str) -> str:
    """Strip light Markdown so text and Word exports read naturally."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?m)^#{1,6}\s*", "", text)
    text = text.replace("`", "")
    return _XML_BAD.sub("", text)


def conversation_markdown(memory: ConversationMemory, title: str = TITLE) -> str:
    return memory.to_markdown(title)


def conversation_text(memory: ConversationMemory, title: str = TITLE) -> str:
    return "\n".join(_plain(line) for line in memory.to_markdown(title).splitlines()) + "\n"


def answer_markdown(turn: Turn, number: int) -> str:
    """One question and answer with its sources."""
    lines = [f"## Question {number}", "", turn.question, "", "**Answer**", "", turn.answer, ""]
    if turn.note:
        lines += [f"_{turn.note}_", ""]
    for warning in turn.warnings:
        lines += [f"> Warning: {warning}", ""]
    if turn.citations:
        lines.append("**Sources**")
        lines += [f"- [{c.number}] {c.label}" for c in turn.citations]
    return "\n".join(lines).strip() + "\n"


def _add_answer_lines(doc: Document, text: str) -> None:
    """Add Markdown-ish answer text to a Word document (bullets become Word bullets)."""
    for raw in text.splitlines():
        line = raw.strip()
        if not line or re.fullmatch(r"\|?[\s\-:|]+\|?", line):  # blank or table separator row
            continue
        if line.startswith(("- ", "* ")):
            doc.add_paragraph(_plain(line[2:]), style="List Bullet")
        elif line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            doc.add_paragraph(_plain("  |  ".join(cells)))
        else:
            doc.add_paragraph(_plain(line))


def conversation_docx(memory: ConversationMemory, title: str = TITLE) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Exported {datetime.now():%Y-%m-%d %H:%M}")
    for number, turn in enumerate(memory.turns, start=1):
        doc.add_heading(f"Question {number}", level=2)
        doc.add_paragraph().add_run(_plain(turn.question)).bold = True
        doc.add_heading("Answer", level=3)
        _add_answer_lines(doc, turn.answer)
        if turn.note:
            doc.add_paragraph().add_run(_plain(turn.note)).italic = True
        for warning in turn.warnings:
            doc.add_paragraph(_plain(f"Warning: {warning}"))
        if turn.citations:
            doc.add_heading("Sources", level=3)
            for citation in turn.citations:
                suffix = "" if citation.cited else " (consulted, not cited)"
                doc.add_paragraph(_plain(f"[{citation.number}] {citation.label}{suffix}"), style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()