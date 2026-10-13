"""PowerPoint (.pptx) loader: slide text, tables, speaker notes, OCR on pictures."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Iterator

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.config import settings
from src.loaders.base import (
    DocElement,
    LoadedDocument,
    LoaderError,
    OcrBudget,
    ProgressFn,
    report,
    rows_to_markdown,
)
from src.ocr import ocr_image
from src.utils import get_logger

log = get_logger(__name__)

MIN_IMAGE_PX = 200
MIN_IMAGE_OCR_CHARS = 10


def _shape_type(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def _walk(shapes) -> Iterator:
    """Yield every shape, going inside groups."""
    for shape in shapes:
        if _shape_type(shape) == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _reading_order(shape) -> tuple[int, int]:
    return (getattr(shape, "top", 0) or 0, getattr(shape, "left", 0) or 0)


def load_pptx(path: Path, display_name: str, progress: ProgressFn | None = None) -> LoadedDocument:
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise LoaderError(
            "This PowerPoint file could not be opened. It may be damaged or password-protected."
        ) from exc

    slides = list(presentation.slides)
    total = len(slides)
    if total > settings.max_pages:
        raise LoaderError(f"This presentation has {total} slides; the limit is {settings.max_pages}.")

    doc = LoadedDocument(filename=display_name, unit_count=total, unit_name="slides")
    budget = OcrBudget()

    for index, slide in enumerate(slides):
        number = index + 1
        location = f"Slide {number}"
        report(progress, index / max(total, 1), f"Reading slide {number} of {total}")
        texts: list[str] = []

        for shape in sorted(_walk(slide.shapes), key=_reading_order):
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    texts.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                markdown = rows_to_markdown(rows)
                if markdown:
                    doc.elements.append(DocElement(markdown, display_name, f"{location} (table)", "table"))
            elif _shape_type(shape) == MSO_SHAPE_TYPE.PICTURE and settings.ocr_enabled:
                try:
                    image = Image.open(io.BytesIO(shape.image.blob))
                    image.load()
                except Exception:
                    continue
                if min(image.size) < MIN_IMAGE_PX:
                    continue
                if not budget.take():
                    continue
                ocr_text = ocr_image(image)
                if len(ocr_text) >= MIN_IMAGE_OCR_CHARS:
                    doc.elements.append(DocElement(ocr_text, display_name, f"{location} (image)", "ocr"))

        if texts:
            doc.elements.append(DocElement("\n".join(texts), display_name, location))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                doc.elements.append(DocElement(notes, display_name, f"{location} (speaker notes)"))

    if budget.used >= budget.limit:
        doc.warnings.append("OCR limit reached: some pictures in this presentation were not read.")
    report(progress, 1.0, "Done")
    return doc