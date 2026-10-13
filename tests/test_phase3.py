"""Phase 3 check: build a sample file of every type and load it."""
from __future__ import annotations

import io
import tempfile
from pathlib import Path

import fitz
import pandas as pd
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from src.loaders import LoaderError, load_document
from src.utils import setup_logging


def text_image(text: str) -> Image.Image:
    """A white image with big black text (so OCR has something to read)."""
    try:
        font = ImageFont.truetype("arial.ttf", 56)
    except OSError:
        font = ImageFont.load_default(size=56)
    image = Image.new("RGB", (1400, 200), "white")
    ImageDraw.Draw(image).text((30, 60), text, fill="black", font=font)
    return image


def make_samples(folder: Path) -> list[Path]:
    paths: list[Path] = []

    txt = folder / "notes.txt"
    txt.write_text("Project Falcon kickoff.\nBudget approved: 50,000 EUR.\n\nOwner: Maria Lopez.", encoding="utf-8")
    paths.append(txt)

    frame = pd.DataFrame({"Item": ["Pen", "Book", "Lamp"], "Price": [1.5, 12, 30], "Stock": [100, 20, 5]})
    csv = folder / "prices.csv"
    frame.to_csv(csv, index=False)
    paths.append(csv)
    xlsx = folder / "sales.xlsx"
    frame.to_excel(xlsx, sheet_name="Sales", index=False)
    paths.append(xlsx)

    word = Document()
    word.add_heading("Project Plan", level=1)
    word.add_paragraph("The project deadline is 30 June 2025.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Phase", "Owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "Design", "Anna"
    docx = folder / "plan.docx"
    word.save(docx)
    paths.append(docx)

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Results"
    slide.placeholders[1].text = "Revenue grew 12 percent in Q3"
    slide.notes_slide.notes_text_frame.text = "Mention the new office"
    slide2 = deck.slides.add_slide(deck.slide_layouts[5])
    slide2.shapes.title.text = "Targets"
    grid = slide2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    grid.cell(0, 0).text, grid.cell(0, 1).text = "Region", "Target"
    grid.cell(1, 0).text, grid.cell(1, 1).text = "North", "900"
    pptx = folder / "deck.pptx"
    deck.save(pptx)
    paths.append(pptx)

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Invoice number INV-2024-001. Total amount due: $1,250.00", fontsize=12)
    text_pdf = folder / "invoice.pdf"
    pdf.save(text_pdf)
    pdf.close()
    paths.append(text_pdf)

    buffer = io.BytesIO()
    text_image("Meeting on 15 March 2025 in Berlin").save(buffer, format="PNG")
    scanned = fitz.open()
    page = scanned.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(36, 36, 559, 36 + 523 * 200 / 1400), stream=buffer.getvalue())
    scanned_pdf = folder / "scanned.pdf"
    scanned.save(scanned_pdf)
    scanned.close()
    paths.append(scanned_pdf)

    png = folder / "photo.png"
    text_image("Order number 77-ALPHA ships Friday").save(png)
    paths.append(png)

    locked = fitz.open()
    locked.new_page().insert_text((72, 72), "secret", fontsize=12)
    locked_pdf = folder / "locked.pdf"
    locked.save(locked_pdf, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw2")
    locked.close()
    paths.append(locked_pdf)

    old = folder / "old.ppt"
    old.write_bytes(b"legacy")
    paths.append(old)
    return paths


def main() -> None:
    setup_logging()
    with tempfile.TemporaryDirectory() as tmp:
        for path in make_samples(Path(tmp)):
            print(f"\n=== {path.name}")
            try:
                doc = load_document(path)
            except LoaderError as exc:
                print("  REJECTED:", exc)
                continue
            print(f"  {len(doc.elements)} piece(s), {doc.unit_count} {doc.unit_name}")
            for element in doc.elements:
                print(f"  [{element.kind}] {element.location}: {element.text[:70]!r}")
            for warning in doc.warnings:
                print("  WARNING:", warning)


if __name__ == "__main__":
    main()